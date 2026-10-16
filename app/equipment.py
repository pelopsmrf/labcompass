"""Учёт оборудования: сроки, статусы, импорт из Excel и выгрузка графика."""
import calendar
import io
import re
from datetime import date, datetime

KINDS = ["СИ", "ИО", "ВО"]
KIND_NAMES = {"СИ": "Средство измерений", "ИО": "Испытательное оборудование", "ВО": "Вспомогательное оборудование"}
SERVICES = ["поверка", "калибровка", "аттестация", "ТО", "не требуется"]
SOON_DAYS = 30  # «скоро срок» — за столько дней
STATUS_NAMES = {"overdue": "Просрочено", "soon": "Скоро срок", "ok": "В порядке", "none": "Срок не указан"}


def add_months(d, months):
    m = d.month - 1 + months
    y, m = d.year + m // 12, m % 12 + 1
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


def parse_date(s):
    """ГГГГ-ММ-ДД, ДД.ММ.ГГГГ, ДД.ММ.ГГ или дата из Excel → date; пусто → None; иначе ValueError."""
    if isinstance(s, datetime):
        return s.date()
    if isinstance(s, date):
        return s
    s = str(s or "").strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"не понимаю дату «{s}»")


def parse_interval(s):
    """Интервал в месяцах: «12», «12 мес», «1 год», «2 года»."""
    s = str(s or "").strip().lower().replace(",", ".")
    if not s:
        return None
    m = re.match(r"^(\d+(?:\.\d+)?)\s*(г|л)?", s)
    if not m:
        raise ValueError(f"не понимаю интервал «{s}»")
    n = float(m.group(1))
    return int(round(n * 12 if m.group(2) else n))


def next_due(last, interval, explicit=None):
    """Срок следующего обслуживания: указанный явно или «последняя дата + интервал»."""
    if explicit:
        return explicit
    if last and interval:
        return add_months(last, interval)
    return None


def status(next_date, today=None):
    if not next_date:
        return "none", None
    today = today or date.today()
    d = date.fromisoformat(next_date) if isinstance(next_date, str) else next_date
    left = (d - today).days
    if left < 0:
        return "overdue", left
    if left <= SOON_DAYS:
        return "soon", left
    return "ok", left


def with_status(rows, today=None):
    out = []
    for r in rows:
        item = dict(r)
        item["status"], item["days_left"] = status(item["next_date"], today)
        out.append(item)
    return out


def norm_kind(s):
    s = str(s or "").strip().lower()
    if s.startswith("и") or "испыт" in s:
        return "ИО"
    if s.startswith("в") or "вспом" in s:
        return "ВО"
    return "СИ"


def norm_service(s):
    s = str(s or "").strip().lower()
    for name in SERVICES:
        if s.startswith(name[:4].lower()):
            return name
    if s.startswith("тех") or s == "то":
        return "ТО"
    if s.startswith("не") or s == "-":
        return "не требуется"
    return "поверка"


IMPORT_COLUMNS = ["Наименование, тип", "Вид (СИ/ИО/ВО)", "Заводской №", "Инвентарный №",
                  "Обслуживание", "Дата последней", "Интервал, мес", "Срок следующей", "№ свидетельства / Аршин"]


def parse_paste(text):
    """Строки, скопированные из Excel (через табуляцию) → (список записей, список ошибок).

    Порядок столбцов — IMPORT_COLUMNS. Строку заголовка пропускаем.
    """
    items, errors = [], []
    for n, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            continue
        cells = [c.strip() for c in line.split("\t")] + [""] * 9
        name = cells[0]
        if not name or name.lower().startswith("наименование"):
            continue
        try:
            last = parse_date(cells[5])
            interval = parse_interval(cells[6])
            nxt = next_due(last, interval, parse_date(cells[7]))
        except ValueError as e:
            errors.append(f"строка {n}: {e}")
            continue
        items.append({
            "name": name[:300], "kind": norm_kind(cells[1]), "serial": cells[2][:100], "inv_no": cells[3][:100],
            "service": norm_service(cells[4]), "last_date": last.isoformat() if last else None,
            "interval_months": interval, "next_date": nxt.isoformat() if nxt else None, "doc_no": cells[8][:200],
        })
    return items, errors


def ru(d):
    return date.fromisoformat(d).strftime("%d.%m.%Y") if d else ""


def to_xlsx(items, org="", today=None):
    """График поверки/калибровки/аттестации в Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    today = today or date.today()
    wb = Workbook()
    ws = wb.active
    ws.title = "График"
    ws["A1"] = f"График метрологического обслуживания оборудования{' — ' + org if org else ''}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Сформирован {today.strftime('%d.%m.%Y')} в ЛабКомпас (labcompass.ru)"
    head = ["№", "Наименование, тип", "Вид", "Заводской №", "Инв. №", "Обслуживание",
            "Дата последней", "Интервал, мес", "Срок следующей", "Осталось дней", "Статус",
            "№ свидетельства / Аршин", "Место / ответственный", "Примечание"]
    ws.append([])
    ws.append(head)
    thin = Side(style="thin", color="BBBBBB")
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF4")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    fills = {"overdue": "F6D5D5", "soon": "FBE8C8"}
    for i, it in enumerate(items, 1):
        ws.append([i, it["name"], it["kind"], it["serial"], it["inv_no"], it["service"], ru(it["last_date"]),
                   it["interval_months"], ru(it["next_date"]), it["days_left"], STATUS_NAMES[it["status"]],
                   it["doc_no"], it["place"], it["note"]])
        if it["status"] in fills:
            ws.cell(ws.max_row, 11).fill = PatternFill("solid", fgColor=fills[it["status"]])
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for col, w in zip("ABCDEFGHIJKLMN", [5, 40, 6, 14, 12, 13, 13, 10, 13, 10, 15, 24, 22, 24]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:N{ws.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
