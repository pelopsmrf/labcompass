"""Реактивы и стандартные образцы (склад) и журнал приготовления растворов: сроки годности, статусы,
вставка из Excel, выгрузка в Excel."""
import io
from datetime import date, timedelta

import equipment as eq

KINDS = ["реактив", "СО", "другое"]
KIND_NAMES = {"реактив": "Реактив", "СО": "Стандартный образец (ГСО, МСО, ОСО, СОП)", "другое": "Другое (среды, газы и т. п.)"}
SOON_DAYS = 30            # реактивы и СО: «скоро срок» — за столько дней
SOLUTION_SOON_DAYS = 2    # растворы: «истекает» — сегодня или в ближайшие 2 дня
STATUS_NAMES = {"overdue": "Просрочен", "soon": "Скоро срок", "ok": "Годен", "none": "Срок не указан"}
SOLUTION_STATUS = {"overdue": "Просрочен", "soon": "Истекает", "ok": "Годен", "none": "Срок не указан"}

FIELDS = ["kind", "name", "grade", "standard", "maker", "lot", "cert_value", "made_date", "shelf_months",
          "expiry_date", "received_date", "opened_date", "qty", "place", "ext_until", "ext_doc", "note"]
SOLUTION_FIELDS = ["name", "conc", "volume", "source_id", "source_text", "method", "prep_date", "shelf_days",
                   "expiry_date", "performer", "note"]
DATES = ["made_date", "expiry_date", "received_date", "opened_date", "ext_until"]


def norm_kind(s):
    s = str(s or "").strip().lower()
    if s in ("со", "гсо", "мсо", "осо", "соп") or s.startswith("станд") or "образ" in s:
        return "СО"
    if s.startswith("др") or s.startswith("сред") or s.startswith("газ"):
        return "другое"
    return "реактив"


def iso(d):
    return d.isoformat() if d else None


def valid_until(item):
    """Срок годности с учётом продления по результатам проверки пригодности."""
    return item.get("ext_until") or item.get("expiry_date")


def status(until, today=None, soon=SOON_DAYS):
    if not until:
        return "none", None
    d = date.fromisoformat(until) if isinstance(until, str) else until
    left = (d - (today or date.today())).days
    return ("overdue" if left < 0 else "soon" if left <= soon else "ok"), left


def with_status(rows, today=None):
    out = []
    for r in rows:
        it = dict(r)
        it["until"] = valid_until(it)
        it["status"], it["days_left"] = status(it["until"], today)
        out.append(it)
    return out


def item_from_form(form):
    """Форма карточки реактива/СО → (запись, ошибка)."""
    it = {f: str(form.get(f, "")).strip() for f in FIELDS}
    limits = {"name": 300, "grade": 50, "standard": 200, "maker": 200, "lot": 100, "cert_value": 200, "qty": 100,
              "place": 200, "ext_doc": 200, "note": 500}
    for f, n in limits.items():
        it[f] = it[f][:n]
    it["kind"] = it["kind"] if it["kind"] in KINDS else "реактив"
    if not it["name"]:
        return it, "Укажите наименование"
    try:
        for f in DATES:
            it[f] = iso(eq.parse_date(it[f]))
        it["shelf_months"] = eq.parse_interval(it["shelf_months"])
    except ValueError as e:
        return it, str(e).capitalize()
    if not it["expiry_date"] and it["made_date"] and it["shelf_months"]:
        it["expiry_date"] = iso(eq.add_months(date.fromisoformat(it["made_date"]), it["shelf_months"]))
    if it["ext_until"] and not it["ext_doc"]:
        return it, "Для продления срока укажите основание — № протокола проверки пригодности"
    return it, None


IMPORT_COLUMNS = ["Наименование", "Вид (реактив/СО/другое)", "Квалификация", "ГОСТ / ТУ / № ГСО", "Изготовитель",
                  "№ партии", "Дата изготовления", "Годен до", "Количество", "Место хранения", "Аттестованное значение"]


def parse_paste(text):
    """Строки из Excel (через табуляцию) → (записи, ошибки). Строку заголовка пропускаем."""
    items, errors = [], []
    for n, line in enumerate(text.replace("\r", "").split("\n"), 1):
        if not line.strip():
            continue
        c = [x.strip() for x in line.split("\t")] + [""] * 11
        if not c[0] or c[0].lower().startswith("наименование"):
            continue
        try:
            made, until = eq.parse_date(c[6]), eq.parse_date(c[7])
        except ValueError as e:
            errors.append(f"строка {n}: {e}")
            continue
        items.append({"kind": norm_kind(c[1]), "name": c[0][:300], "grade": c[2][:50], "standard": c[3][:200],
                      "maker": c[4][:200], "lot": c[5][:100], "made_date": iso(made), "expiry_date": iso(until),
                      "qty": c[8][:100], "place": c[9][:200], "cert_value": c[10][:200]})
    return items, errors


# ---------- растворы ----------

def solution_from_form(form, sources):
    """Форма записи журнала → (запись, ошибка). sources — {id: запись склада} этой лаборатории."""
    it = {f: str(form.get(f, "")).strip() for f in SOLUTION_FIELDS}
    for f, n in {"name": 300, "conc": 100, "volume": 50, "source_text": 300, "method": 300, "performer": 100,
                 "note": 500}.items():
        it[f] = it[f][:n]
    if not it["name"]:
        return it, "Укажите, что приготовлено"
    src = it["source_id"]
    it["source_id"] = int(src) if src.isdigit() and int(src) in sources else None
    try:
        prep = eq.parse_date(it["prep_date"])
        explicit = eq.parse_date(it["expiry_date"])
    except ValueError as e:
        return it, str(e).capitalize()
    if not prep:
        return it, "Укажите дату приготовления"
    days = it["shelf_days"]
    if days and not (days.isdigit() and 0 < int(days) <= 3650):
        return it, "Срок хранения — целое число дней"
    it["shelf_days"] = int(days) if days else None
    it["prep_date"] = prep.isoformat()
    if it["shelf_days"]:
        explicit = prep + timedelta(days=it["shelf_days"])
    it["expiry_date"] = iso(explicit)
    return it, None


def source_problem(sol, src):
    """Исходный реактив/СО на дату приготовления был просрочен → текст предупреждения."""
    if not src:
        return None
    until = valid_until(src)
    if until and sol["prep_date"] and until < sol["prep_date"]:
        return f"исходный {'СО' if src['kind'] == 'СО' else 'реактив'} на дату приготовления был просрочен (годен до {eq.ru(until)})"
    return None


def solutions_with_status(rows, sources, today=None):
    out = []
    for r in rows:
        it = dict(r)
        it["status"], it["days_left"] = status(it["expiry_date"], today, SOLUTION_SOON_DAYS)
        src = sources.get(it["source_id"])
        it["source"] = src
        it["source_label"] = source_label(src) if src else it["source_text"]
        it["problem"] = source_problem(it, src)
        out.append(it)
    return out


def source_label(src):
    parts = [src["name"]]
    if src["grade"]:
        parts.append(src["grade"])
    if src["standard"]:
        parts.append(src["standard"])
    if src["lot"]:
        parts.append("партия " + src["lot"])
    return ", ".join(parts)


# ---------- Excel ----------

def _sheet(title, sub, head, rows, widths, fills_col=None, statuses=()):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = sub
    ws.append([])
    ws.append(head)
    thin = Side(style="thin", color="BBBBBB")
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF4")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    fills = {"overdue": "F6D5D5", "soon": "FBE8C8"}
    for row, st in zip(rows, statuses or [None] * len(rows)):
        ws.append(row)
        if fills_col and st in fills:
            ws.cell(ws.max_row, fills_col).fill = PatternFill("solid", fgColor=fills[st])
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{chr(64 + len(head))}{ws.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def stock_xlsx(items, org="", today=None):
    today = today or date.today()
    head = ["№", "Наименование", "Вид", "Квалификация", "ГОСТ / ТУ / № ГСО", "Изготовитель", "№ партии",
            "Аттестованное значение", "Дата изготовления", "Годен до", "Продлён до", "Основание продления",
            "Статус", "Получен", "Вскрыт", "Количество", "Место хранения", "Примечание"]
    rows = [[i, it["name"], it["kind"], it["grade"], it["standard"], it["maker"], it["lot"], it["cert_value"],
             eq.ru(it["made_date"]), eq.ru(it["expiry_date"]), eq.ru(it["ext_until"]), it["ext_doc"],
             STATUS_NAMES[it["status"]], eq.ru(it["received_date"]), eq.ru(it["opened_date"]), it["qty"],
             it["place"], it["note"]] for i, it in enumerate(items, 1)]
    return _sheet(f"Реактивы и стандартные образцы{' — ' + org if org else ''}",
                  f"Сформирован {today.strftime('%d.%m.%Y')} в ЛабКомпас (labcompass.ru)", head, rows,
                  [5, 36, 9, 10, 20, 20, 12, 18, 12, 12, 12, 18, 12, 11, 11, 12, 18, 22], 13,
                  [it["status"] for it in items])


def solutions_xlsx(items, org="", today=None):
    today = today or date.today()
    head = ["№", "Дата приготовления", "Наименование раствора", "Концентрация", "Объём", "Из чего приготовлен",
            "Методика (НД)", "Срок хранения, сут", "Годен до", "Статус", "Приготовил", "Примечание"]
    rows = [[i, eq.ru(it["prep_date"]), it["name"], it["conc"], it["volume"], it["source_label"], it["method"],
             it["shelf_days"], eq.ru(it["expiry_date"]), SOLUTION_STATUS[it["status"]], it["performer"],
             "; ".join(x for x in (it["note"], it["problem"] and "ВНИМАНИЕ: " + it["problem"]) if x)]
            for i, it in enumerate(items, 1)]
    return _sheet(f"Журнал приготовления растворов{' — ' + org if org else ''}",
                  f"Сформирован {today.strftime('%d.%m.%Y')} в ЛабКомпас (labcompass.ru)", head, rows,
                  [5, 12, 34, 14, 9, 36, 24, 10, 12, 11, 16, 28], 10, [it["status"] for it in items])
